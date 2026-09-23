from io import BytesIO

from docx import Document
import pytest


def document_text(content):
    document = Document(BytesIO(content))
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def risk_report_payload():
    return {
        "project_name": "风险报告脱敏项目",
        "tender_company": "脱敏招标单位",
        "deadline": "2026-09-30",
        "deposit": "10万元",
        "procurement_requirements": [
            {
                "item_name": "脱敏设备",
                "specification": "规格甲",
                "quantity": "2台",
                "budget": "10万元",
                "import_allowed": "否",
                "page": 2,
            }
        ],
        "risk": [
            {
                "level": "高风险",
                "page": 3,
                "keyword": "资格审查",
                "reason": "材料缺失",
                "suggestion": "补齐材料",
                "deduction": 25,
            },
            {"level": "中风险", "page": 5, "deduction": 8},
            {"level": "低风险", "page": 8, "deduction": 3},
        ],
    }


def compare_report_payload():
    return {
        "projectA": {
            "project_name": "对比项目甲",
            "score": 88,
            "score_level": "低风险",
            "risk_count": 2,
            "high_count": 1,
            "middle_count": 1,
            "low_count": 0,
            "total_deduction": 12,
            "analysis": {"tender_company": "甲单位"},
            "risk": [],
        },
        "projectB": {
            "project_name": "对比项目乙",
            "score": 76,
            "score_level": "中低风险",
            "risk_count": 3,
            "high_count": 1,
            "middle_count": 1,
            "low_count": 1,
            "total_deduction": 24,
            "analysis": {"tender_company": "乙单位"},
            "risk": [],
        },
        "decisionScoreA": 83,
        "decisionScoreB": 71,
        "recommended": "A",
        "confidenceLevel": "高",
        "commonCategories": [],
        "riskStrengthCompare": [],
        "onlyARisks": [],
        "onlyBRisks": [],
    }


def test_risk_report_matches_canonical_page_fields(client, isolated_app):
    first = client.post("/api/report", json=risk_report_payload())
    second = client.post("/api/report", json=risk_report_payload())

    assert first.status_code == 200
    assert second.status_code == 200
    text = document_text(first.content)
    assert "项目名称：风险报告脱敏项目" in text
    assert "投标建议指数：64分" in text
    assert "风险评级：中风险" in text
    assert "风险统计：共 3项，高风险 1项，中风险 1项，低风险 1项" in text
    assert "风险总扣分：-36分" in text
    assert "资格审查" in text
    assert "五、采购清单摘要" in text
    assert "标的：脱敏设备" in text
    assert "数量：2台" in text
    reports = list((isolated_app["runtime_dir"] / "reports").glob("*.docx"))
    assert len(reports) == 2
    assert len({path.name for path in reports}) == 2


def test_compare_report_matches_page_project_fields(client):
    response = client.post("/api/compare-report", json=compare_report_payload())

    assert response.status_code == 200
    text = document_text(response.content)
    assert "对比项目甲" in text
    assert "对比项目乙" in text
    assert "88 分" in text
    assert "76 分" in text
    assert "低风险" in text
    assert "中低风险" in text
    assert "1 项" in text
    assert "-12 分" in text
    assert "-24 分" in text


def test_risk_report_prefers_rendered_page_mapping(client):
    payload = {
        "project_name": "Word映射项目",
        "risk": [
            {
                "level": "中风险",
                "page": 1,
                "report_page": 2,
                "keyword": "渲染页",
                "reason": "页码映射",
                "suggestion": "核对预览",
                "deduction": 8,
            }
        ],
    }

    response = client.post("/api/report", json=payload)

    assert response.status_code == 200
    text = document_text(response.content)
    assert "第2页" in text
    assert "第1页" not in text


def test_reports_open_with_missing_fields(client):
    risk_response = client.post("/api/report", json={})
    compare_response = client.post("/api/compare-report", json={})

    assert risk_response.status_code == 200
    assert compare_response.status_code == 200
    assert "项目名称：未找到" in document_text(risk_response.content)
    assert "未命名项目" in document_text(compare_response.content)


@pytest.mark.parametrize("procurement", [None, [], "invalid", [{"item_name": "合成设备"}]])
def test_risk_report_section_numbers_are_contiguous(client, procurement):
    payload = risk_report_payload()
    if procurement is None:
        payload.pop("procurement_requirements")
    else:
        payload["procurement_requirements"] = procurement
    response = client.post("/api/report", json=payload)
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 2"]
    titles = ["项目基本信息", "投标风险评估", "技术要求摘要", "商务资格要求"]
    if isinstance(procurement, list) and procurement:
        titles.append("采购清单摘要")
    titles.extend(["风险明细", "风险原文依据", "AI投标建议"])
    assert headings == [f"{number}、{title}" for number, title in zip("一二三四五六七八", titles)]
