import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.config import get_report_dir


# =========================================================
# 基础工具
# =========================================================

def safe(value, default="未找到"):

    if value is None:
        return default

    value = str(value).strip()

    if not value:
        return default

    return value


def get_project_name(project):

    return (
        project.get("project_name")
        or project.get("filename")
        or "未命名项目"
    )


def set_cell_background(cell, color):

    tc_pr = cell._tc.get_or_add_tcPr()

    # 先删除旧背景，避免重复叠加
    for child in tc_pr.findall(
        qn("w:shd")
    ):
        tc_pr.remove(child)

    shd = OxmlElement("w:shd")

    shd.set(
        qn("w:val"),
        "clear"
    )

    shd.set(
        qn("w:color"),
        "auto"
    )

    shd.set(
        qn("w:fill"),
        color
    )

    tc_pr.append(shd)


def add_field(paragraph, field_code):

    run = paragraph.add_run()

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(
        qn("w:fldCharType"),
        "begin"
    )

    instr_text = OxmlElement("w:instrText")
    instr_text.set(
        qn("xml:space"),
        "preserve"
    )
    instr_text.text = field_code

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(
        qn("w:fldCharType"),
        "separate"
    )

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(
        qn("w:fldCharType"),
        "end"
    )

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)


def add_page_number(paragraph):

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "第 "
    )

    run.font.size = Pt(8)

    add_field(
        paragraph,
        "PAGE"
    )

    run = paragraph.add_run(
        " 页 / 共 "
    )

    run.font.size = Pt(8)

    add_field(
        paragraph,
        "NUMPAGES"
    )

    run = paragraph.add_run(
        " 页"
    )

    run.font.size = Pt(8)

def set_cell_text(
    cell,
    text,
    bold=False,
    size=10,
    color=None
):

    cell.text = ""

    paragraph = cell.paragraphs[0]

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        str(text)
    )

    run.bold = bold
    run.font.size = Pt(size)

    run.font.name = "微软雅黑"

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "微软雅黑"
    )

    if color:

        run.font.color.rgb = RGBColor(
            *color
        )

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )

def get_risk_color(level):

    level = str(level or "")

    if "高" in level:
        return "FDE2E1"

    if "中" in level:
        return "FCE8D5"

    if "低" in level:
        return "E5F5E0"

    return "F2F2F2"


def add_info_row(
    table,
    label,
    value_a,
    value_b
):

    cells = table.add_row().cells

    set_cell_text(
        cells[0],
        label,
        bold=True
    )

    set_cell_text(
        cells[1],
        value_a
    )

    set_cell_text(
        cells[2],
        value_b
    )


def add_section_title(
    doc,
    number,
    title
):

    paragraph = doc.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(8)

    run = paragraph.add_run(
        f"{number}  {title}"
    )

    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "微软雅黑"

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "微软雅黑"
    )

    run.font.color.rgb = RGBColor(
        31,
        78,
        121
    )


def add_risk_item(
    doc,
    risk,
    index
):

    level = safe(
        risk.get("level"),
        "未知"
    )

    category = safe(
        risk.get("category"),
        "其他风险"
    )

    keyword = safe(
        risk.get("keyword"),
        "未提供"
    )

    page = safe(
        risk.get("page"),
        "-"
    )

    reason = safe(
        risk.get("reason"),
        ""
    )

    suggestion = safe(
        risk.get("suggestion"),
        ""
    )

    deduction = (
        risk.get("deduction")
        or 0
    )


    p = doc.add_paragraph()

    run = p.add_run(
        f"{index}. [{level}风险] "
        f"{category} - {keyword}"
    )

    run.bold = True


    doc.add_paragraph(
        f"所在页码：第 {page} 页"
    )

    doc.add_paragraph(
        f"风险扣分：-{deduction} 分"
    )


    if reason != "未找到":

        doc.add_paragraph(
            f"风险原因：{reason}"
        )


    if suggestion != "未找到":

        doc.add_paragraph(
            f"处理建议：{suggestion}"
        )


# =========================================================
# 创建报告
# =========================================================

def create_compare_report(data):

    doc = Document()

    ai_decision = data.get(
        "aiDecision",
        {}
    ) or {}

    project_a = data.get(
        "projectA",
        {}
    ) or {}


    # =====================================================
    # 页面设置
    # =====================================================

    section = doc.sections[0]

    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)


    # =====================================================
    # 默认字体
    # =====================================================

    normal_style = doc.styles["Normal"]

    normal_style.font.name = "微软雅黑"
    normal_style.font.size = Pt(10.5)

    normal_style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "微软雅黑"
    )


    project_a = data.get(
        "projectA",
        {}
    ) or {}

    project_b = data.get(
        "projectB",
        {}
    ) or {}


    analysis_a = project_a.get(
        "analysis",
        {}
    ) or {}

    analysis_b = project_b.get(
        "analysis",
        {}
    ) or {}


    name_a = get_project_name(
        project_a
    )

    name_b = get_project_name(
        project_b
    )


    # =====================================================
    # 页眉
    # =====================================================

    header = section.header

    p = header.paragraphs[0]

    p.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    run = p.add_run(
        "AI Bid Assistant ｜ 投标决策分析"
    )

    run.font.size = Pt(9)

    run.font.color.rgb = RGBColor(
        120,
        120,
        120
    )


# =====================================================
# 页脚
# =====================================================

    # =====================================================
    # 页脚
    # =====================================================

    footer = section.footer


    # 第一行：报告说明
    p = footer.paragraphs[0]

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = p.add_run(
        "AI生成内容仅作为投标决策辅助，请结合招标文件原文进行最终审核"
    )

    run.font.size = Pt(8)

    run.font.color.rgb = RGBColor(
        130,
        130,
        130
    )


    # 第二行：自动页码
    page_p = footer.add_paragraph()

    add_page_number(
        page_p
    )


    # =====================================================
    # 封面
    # =====================================================

    for _ in range(5):

        doc.add_paragraph()


    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        "AI 投标项目对比分析报告"
    )

    run.bold = True
    run.font.size = Pt(26)

    run.font.color.rgb = RGBColor(
        31,
        78,
        121
    )


    subtitle = doc.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle.paragraph_format.space_before = Pt(16)

    run = subtitle.add_run(
        "项目投标决策辅助报告"
    )

    run.font.size = Pt(16)

    run.font.color.rgb = RGBColor(
        100,
        100,
        100
    )


    doc.add_paragraph()


    project_text = doc.add_paragraph()

    project_text.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    project_text.add_run(
        f"\n项目 A：{name_a}"
    ).bold = True

    project_text.add_run(
        f"\n\nVS\n\n项目 B：{name_b}"
    ).bold = True


    doc.add_paragraph()


    date_p = doc.add_paragraph()

    date_p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    date_p.add_run(
        "报告生成时间："
        +
        datetime.now().strftime(
            "%Y-%m-%d %H:%M"
        )
    )


    doc.add_page_break()


        # =====================================================
    # 一、管理层决策摘要
    # =====================================================

    add_section_title(
        doc,
        "一",
        "管理层决策摘要"
    )


    score_a = data.get(
        "decisionScoreA",
        0
    )

    score_b = data.get(
        "decisionScoreB",
        0
    )


    recommended = data.get(
        "recommended",
        "持平"
    )


    confidence_level = safe(
        data.get(
            "confidenceLevel"
        ),
        "未评估"
    )


    score_difference = data.get(
        "scoreDifference",
        0
    )


    if recommended == "A":

        recommended_name = name_a

    elif recommended == "B":

        recommended_name = name_b

    else:

        recommended_name = (
            "两个项目综合表现接近"
        )


    # =====================================================
    # A / B 综合指数
    # =====================================================

    summary_table = doc.add_table(
        rows=2,
        cols=3
    )

    summary_table.style = "Table Grid"

    summary_table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )


    # 第一行

    set_cell_text(
        summary_table.cell(0, 0),
        "核心指标",
        bold=True
    )

    set_cell_text(
        summary_table.cell(0, 1),
        "项目 A",
        bold=True
    )

    set_cell_text(
        summary_table.cell(0, 2),
        "项目 B",
        bold=True
    )


    for i in range(3):

        set_cell_background(
            summary_table.cell(0, i),
            "D9EAF7"
        )


    # 第二行

    set_cell_text(
        summary_table.cell(1, 0),
        "综合指数",
        bold=True,
        size=12
    )


    set_cell_text(
       summary_table.cell(1, 1),
       f"{score_a} 分",
      bold=True,
      size=20
    )

    set_cell_text(
      summary_table.cell(1, 2),
      f"{score_b} 分",
      bold=True,
      size=20
    )


    # 推荐项目增加背景

    if recommended == "A":

        set_cell_background(
          summary_table.cell(1, 1),
         "E8F5E9"
    )

        set_cell_background(
          summary_table.cell(1, 2),
          "F5F5F5"
        )


    elif recommended == "B":

        set_cell_background(
          summary_table.cell(1, 1),
          "F5F5F5"
        )

        set_cell_background(
          summary_table.cell(1, 2),
         "E8F5E9"
        )


    doc.add_paragraph()


    # =====================================================
    # 推荐结论
    # =====================================================

    decision_table = doc.add_table(
        rows=4,
        cols=2
    )

    decision_table.style = "Table Grid"

    decision_table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )


    decision_rows = [

        (
            "推荐项目",
            recommended_name
        ),

        (
            "推荐置信度",
            confidence_level
        ),

        (
            "综合指数差值",
            f"{score_difference} 分"
        ),

        (
            "推荐依据",
            safe(
                data.get(
                    "recommendationReason"
                ),
                "暂无"
            )
        )

    ]


    for row_index, (
        label,
        value

    ) in enumerate(decision_rows):

        set_cell_text(
            decision_table.cell(row_index, 0),
            label,
            bold=True,
        )

        if label == "推荐项目":

            set_cell_background(
                decision_table.cell(row_index, 1),
                "E8F5E9",
            )

            set_cell_background(
                decision_table.cell(row_index, 0),
                "F2F2F2",
            )

            set_cell_text(
                decision_table.cell(row_index, 1),
                value,
            )

    doc.add_paragraph()


    # =====================================================
    # 风险核心指标
    # =====================================================

    p = doc.add_paragraph()

    run = p.add_run(
        "关键风险指标"
    )

    run.bold = True
    run.font.size = Pt(13)


    risk_table = doc.add_table(
        rows=1,
        cols=3
    )

    risk_table.style = "Table Grid"

    risk_table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )


    risk_headers = [
        "风险指标",
        "项目 A",
        "项目 B"
    ]


    for i, value in enumerate(
        risk_headers
    ):

        set_cell_text(
            risk_table.rows[0].cells[i],
            value,
            bold=True
        )

        set_cell_background(
            risk_table.rows[0].cells[i],
            "EAF2F8"
        )


    risk_rows = [

        (
            "高风险",
            project_a.get(
                "high_count",
                0
            ),
            project_b.get(
                "high_count",
                0
            )
        ),

        (
            "中风险",
            project_a.get(
                "middle_count",
                0
            ),
            project_b.get(
                "middle_count",
                0
            )
        ),

        (
            "低风险",
            project_a.get(
                "low_count",
                0
            ),
            project_b.get(
                "low_count",
                0
            )
        ),

        (
            "风险总扣分",
            f"-{project_a.get('total_deduction', 0)}",
            f"-{project_b.get('total_deduction', 0)}"
        ),

        (
            "综合指数",
            score_a,
            score_b
        )

    ]


    for label, value_a, value_b in risk_rows:

        cells = risk_table.add_row().cells

        set_cell_text(
            cells[0],
            label,
            bold=True
        )

        set_cell_text(
            cells[1],
            value_a
        )

        set_cell_text(
            cells[2],
            value_b
        )

        # =========================
        # 风险颜色
        # =========================
        if label == "高风险":

            set_cell_background(
                cells[0],
                "FF9999"
            )

            set_cell_background(
                cells[1],
                "FF9999"
            )

            set_cell_background(
                cells[2],
                "FF9999"
            )


        elif label == "中风险":

            set_cell_background(
                cells[0],
                "FFD966"
            )

            set_cell_background(
                cells[1],
                "FFD966"
            )

            set_cell_background(
                cells[2],
                "FFD966"
            )


        elif label == "低风险":

            set_cell_background(
                cells[0],
                "C6EFCE"
            )

            set_cell_background(
                cells[1],
                "C6EFCE"
            )

            set_cell_background(
                cells[2],
                "C6EFCE"
            )

        elif label == "风险总扣分":

            set_cell_background(
                cells[1],
                "FFF1F0"
            )

            set_cell_background(
                cells[2],
                "FFF1F0"
            )


        elif label == "综合指数":

            set_cell_background(
                cells[1],
                "E6F4FF"
            )

            set_cell_background(
                cells[2],
                "E6F4FF"
            )


    # =====================================================
    # 置信度说明
    # =====================================================

    confidence_text = data.get(
        "confidenceText"
    )


    if confidence_text:

        doc.add_paragraph()

        p = doc.add_paragraph()

        p.add_run(
            "决策说明："
        ).bold = True

        p.add_run(
            str(confidence_text)
        )


    # =====================================================
    # 决策提醒
    # =====================================================

    doc.add_paragraph()


    p = doc.add_paragraph()

    run = p.add_run(
        "⚠ 决策前重点复核"
    )

    run.bold = True

    run.font.size = Pt(12)


    review_items = [

        "企业资格及资格审查条件",

        "项目负责人及人员证书",

        "投标保证金",

        "废标 / 否决投标条款",

        "电子签名及电子签章",

        "关键工期及履约要求"

    ]


    for item in review_items:

        doc.add_paragraph(
            f"• {item}"
        )

    # =====================================================
    # AI管理层决策摘要
    # =====================================================

    doc.add_paragraph()


    p = doc.add_paragraph()

    run = p.add_run(
        "AI 管理层决策摘要"
    )

    run.bold = True
    run.font.size = Pt(14)

    run.font.color.rgb = RGBColor(
        31,
        78,
        121
    )


    if ai_decision:


        # =========================
        # 执行摘要
        # =========================

        executive_summary = ai_decision.get(
            "executive_summary",
            ""
        )


        if executive_summary:

            p = doc.add_paragraph()

            p.add_run(
                "执行摘要："
            ).bold = True

            p.add_run(
                str(executive_summary)
            )


        # =========================
        # 推荐类型
        # =========================

        recommendation_type = ai_decision.get(
            "recommendation_type",
            ""
        )


        if recommendation_type:

            p = doc.add_paragraph()

            p.add_run(
                "推荐类型："
            ).bold = True

            p.add_run(
                str(recommendation_type)
            )


        # =========================
        # 推荐项目优势
        # =========================

        advantages = ai_decision.get(
            "key_advantages",
            []
        ) or []


        if advantages:

            p = doc.add_paragraph()

            p.add_run(
                "推荐项目主要优势"
            ).bold = True


            for index, item in enumerate(
                advantages,
                1
            ):

                doc.add_paragraph(
                    f"{index}. {item}"
                )


        # =========================
        # 核心风险
        # =========================

        key_risks = ai_decision.get(
            "key_risks",
            []
        ) or []


        if key_risks:

            p = doc.add_paragraph()

            p.add_run(
                "推荐项目仍需关注的核心风险"
            ).bold = True


            for index, item in enumerate(
                key_risks,
                1
            ):

                doc.add_paragraph(
                    f"{index}. {item}"
                )


        # =========================
        # 投标前必须解决
        # =========================

        must_resolve = ai_decision.get(
            "must_resolve_before_bid",
            []
        ) or []


        if must_resolve:

            p = doc.add_paragraph()

            p.add_run(
                "投标前必须处理事项"
            ).bold = True


            for index, item in enumerate(
                must_resolve,
                1
            ):

                doc.add_paragraph(
                    f"{index}. {item}"
                )


        # =========================
        # 改变推荐的条件
        # =========================

        switch_conditions = ai_decision.get(
            "switch_conditions",
            []
        ) or []


        if switch_conditions:

            p = doc.add_paragraph()

            p.add_run(
                "需要重新评估当前推荐的情况"
            ).bold = True


            for index, item in enumerate(
                switch_conditions,
                1
            ):

                doc.add_paragraph(
                    f"{index}. {item}"
                )


        # =========================
        # 管理层最终建议
        # =========================

        management_advice = ai_decision.get(
            "management_advice",
            ""
        )


        if management_advice:

            p = doc.add_paragraph()

            p.add_run(
                "管理层建议："
            ).bold = True

            p.add_run(
                str(management_advice)
            )


    else:

        doc.add_paragraph(
            "本次对比未生成AI决策摘要，"
            "请结合系统评分及风险明细进行人工复核。"
        )


    # =====================================================
    # 二、项目核心信息对比
    # =====================================================

    add_section_title(
        doc,
        "二",
        "项目核心信息对比"
    )


    table = doc.add_table(
        rows=1,
        cols=3
    )

    table.style = "Table Grid"

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )


    headers = [
        "对比项",
        "项目 A",
        "项目 B"
    ]


    for i, value in enumerate(
        headers
    ):

        set_cell_text(
            table.rows[0].cells[i],
            value,
            bold=True
        )

        set_cell_background(
            table.rows[0].cells[i],
            "D9EAF7"
        )


    add_info_row(
        table,
        "项目名称",
        name_a,
        name_b
    )

    add_info_row(
        table,
        "招标单位",
        safe(
            analysis_a.get(
                "tender_company"
            )
        ),
        safe(
            analysis_b.get(
                "tender_company"
            )
        )
    )

    add_info_row(
        table,
        "投标截止时间",
        safe(
            analysis_a.get(
                "deadline"
            )
        ),
        safe(
            analysis_b.get(
                "deadline"
            )
        )
    )

    add_info_row(
        table,
        "保证金",
        safe(
            analysis_a.get(
                "deposit"
            )
        ),
        safe(
            analysis_b.get(
                "deposit"
            )
        )
    )

    add_info_row(
        table,
        "AI风险评分",
        f"{project_a.get('score', 0)} 分",
        f"{project_b.get('score', 0)} 分"
    )

    add_info_row(
        table,
        "风险评级",
        safe(
            project_a.get(
                "score_level"
            )
        ),
        safe(
            project_b.get(
                "score_level"
            )
        )
    )

    add_info_row(
        table,
        "高风险",
        f"{project_a.get('high_count', 0)} 项",
        f"{project_b.get('high_count', 0)} 项"
    )

    add_info_row(
        table,
        "中风险",
        f"{project_a.get('middle_count', 0)} 项",
        f"{project_b.get('middle_count', 0)} 项"
    )

    add_info_row(
        table,
        "低风险",
        f"{project_a.get('low_count', 0)} 项",
        f"{project_b.get('low_count', 0)} 项"
    )

    add_info_row(
        table,
        "风险总扣分",
        f"-{project_a.get('total_deduction', 0)} 分",
        f"-{project_b.get('total_deduction', 0)} 分"
    )


    # =====================================================
    # 三、综合推荐
    # =====================================================

    add_section_title(
        doc,
        "三",
        "综合推荐与决策解释"
    )


    p = doc.add_paragraph()

    p.add_run(
        "推荐项目："
    ).bold = True

    p.add_run(
        recommended_name
    )


    p = doc.add_paragraph()

    p.add_run(
        "推荐置信度："
    ).bold = True

    p.add_run(
        safe(
            data.get(
                "confidenceLevel"
            ),
            "未评估"
        )
    )


    p = doc.add_paragraph()

    p.add_run(
        "系统分析："
    ).bold = True

    p.add_run(
        safe(
            data.get(
                "recommendationReason"
            ),
            "暂无推荐说明"
        )
    )


    # =====================================================
    # 四、评分解释明细
    # =====================================================

    add_section_title(
        doc,
        "四",
        "综合指数评分明细"
    )


    score_table = doc.add_table(
        rows=1,
        cols=3
    )

    score_table.style = "Table Grid"


    for i, text in enumerate([
        "评分项",
        "项目 A",
        "项目 B"
    ]):

        set_cell_text(
            score_table.rows[0].cells[i],
            text,
            bold=True
        )

        set_cell_background(
            score_table.rows[0].cells[i],
            "EAF2F8"
        )


    score_rows = [

        (
            "原始 AI 评分",
            data.get(
                "baseScoreA",
                project_a.get(
                    "score",
                    0
                )
            ),
            data.get(
                "baseScoreB",
                project_b.get(
                    "score",
                    0
                )
            )
        ),

        (
            "高风险惩罚",
            -data.get(
                "highPenaltyA",
                0
            ),
            -data.get(
                "highPenaltyB",
                0
            )
        ),

        (
            "独有高风险惩罚",
            -data.get(
                "uniqueHighPenaltyA",
                0
            ),
            -data.get(
                "uniqueHighPenaltyB",
                0
            )
        ),

        (
            "共同风险强度惩罚",
            -data.get(
                "commonStrengthPenaltyA",
                0
            ),
            -data.get(
                "commonStrengthPenaltyB",
                0
            )
        ),

        (
            "最终综合指数",
            score_a,
            score_b
        )

    ]


    for (
        label,
        value_a,
        value_b
    ) in score_rows:

        cells = (
            score_table
            .add_row()
            .cells
        )

        set_cell_text(
            cells[0],
            label
        )

        set_cell_text(
            cells[1],
            f"{value_a} 分"
        )

        set_cell_text(
            cells[2],
            f"{value_b} 分"
        )


    # =====================================================
    # 五、共同风险
    # =====================================================

    add_section_title(
        doc,
        "五",
        "共同风险分析"
    )


    common_categories = data.get(
        "commonCategories",
        []
    ) or []


    if common_categories:

        doc.add_paragraph(
            "两个项目共同存在以下风险类别："
        )

        for index, category in enumerate(
            common_categories,
            1
        ):

            doc.add_paragraph(
                f"{index}. {category}"
            )

    else:

        doc.add_paragraph(
            "当前未识别到明显的共同风险类别。"
        )


    # =====================================================
    # 六、共同风险强度
    # =====================================================

    add_section_title(
        doc,
        "六",
        "共同风险强度对比"
    )


    strength_list = data.get(
        "riskStrengthCompare",
        []
    ) or []


    if not strength_list:

        doc.add_paragraph(
            "暂无共同风险强度对比数据。"
        )


    for index, item in enumerate(
        strength_list,
        1
    ):

        category = safe(
            item.get(
                "category"
            ),
            "其他风险"
        )


        p = doc.add_paragraph()

        run = p.add_run(
            f"{index}. {category}"
        )

        run.bold = True


        risk_a = item.get(
            "riskA",
            {}
        ) or {}

        risk_b = item.get(
            "riskB",
            {}
        ) or {}


        doc.add_paragraph(
            "项目 A："
            +
            f"{safe(risk_a.get('level'), '未知')}风险；"
            +
            f"扣分 {item.get('deductionA', 0)} 分；"
            +
            f"关键词：{safe(risk_a.get('keyword'), '-')}"
        )


        doc.add_paragraph(
            "项目 B："
            +
            f"{safe(risk_b.get('level'), '未知')}风险；"
            +
            f"扣分 {item.get('deductionB', 0)} 分；"
            +
            f"关键词：{safe(risk_b.get('keyword'), '-')}"
        )


        doc.add_paragraph(
            "对比结论："
            +
            safe(
                item.get(
                    "compareReason"
                ),
                "暂无"
            )
        )


    # =====================================================
    # 七、项目A独有风险
    # =====================================================

    add_section_title(
        doc,
        "七",
        f"项目 A 独有风险｜{name_a}"
    )


    only_a = data.get(
        "onlyARisks",
        []
    ) or []


    if only_a:

        for index, risk in enumerate(
            only_a,
            1
        ):

            add_risk_item(
                doc,
                risk,
                index
            )

    else:

        doc.add_paragraph(
            "未识别到项目 A 独有风险。"
        )


    # =====================================================
    # 八、项目B独有风险
    # =====================================================

    add_section_title(
        doc,
        "八",
        f"项目 B 独有风险｜{name_b}"
    )


    only_b = data.get(
        "onlyBRisks",
        []
    ) or []


    if only_b:

        for index, risk in enumerate(
            only_b,
            1
        ):

            add_risk_item(
                doc,
                risk,
                index
            )

    else:

        doc.add_paragraph(
            "未识别到项目 B 独有风险。"
        )


    # =====================================================
    # 九、最终投标决策建议
    # =====================================================

    add_section_title(
        doc,
        "九",
        "最终投标决策建议"
    )


    if recommended == "A":

        decision_text = (
            f"根据当前招标文件风险分析及项目对比结果，"
            f"项目 A「{name_a}」的综合指数高于项目 B。"
            f"在现有分析条件下，可优先考虑项目 A。"
        )

    elif recommended == "B":

        decision_text = (
            f"根据当前招标文件风险分析及项目对比结果，"
            f"项目 B「{name_b}」的综合指数高于项目 A。"
            f"在现有分析条件下，可优先考虑项目 B。"
        )

    else:

        decision_text = (
            "两个项目综合指数较为接近，"
            "当前风险分析不足以形成明显优先推荐。"
            "建议进一步结合项目利润率、企业资源占用、"
            "中标概率、工期安排及履约能力进行决策。"
        )


    doc.add_paragraph(
        decision_text
    )


    doc.add_paragraph(
        "建议在正式投标前，由商务、技术、法务及项目负责人"
        "再次核验资格条件、废标条款、保证金、电子签章、"
        "人员证书、工期要求及关键合同风险。"
    )


    # =====================================================
    # 十、免责声明
    # =====================================================

    add_section_title(
        doc,
        "十",
        "报告说明"
    )


    doc.add_paragraph(
        "本报告由 AI Bid Assistant 根据上传的招标文件及系统分析结果自动生成。"
        "报告内容仅作为投标风险识别和内部决策辅助，不构成法律、财务或最终投标决策意见。"
        "涉及资格审查、废标条件、合同责任及金额等关键内容，应以招标文件原文及人工复核结果为准。"
    )


    # =====================================================
    # 保存
    # =====================================================

    report_dir = get_report_dir()
    report_dir.mkdir(parents=True, exist_ok=True)


    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S_%f"
    )


    filename = (
        f"AI投标项目对比分析报告_{timestamp}.docx"
    )


    path = os.path.join(
        report_dir,
        filename
    )


    doc.save(
        path
    )


    return os.path.abspath(
        path
    )
