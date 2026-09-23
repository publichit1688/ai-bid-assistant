import os
from datetime import datetime

from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.services.risk_scoring import calculate_risk_summary
from app.config import get_report_dir


def format_procurement_item(item):
    if isinstance(item, str):
        return item
    if not isinstance(item, dict):
        return ""
    labels = (
        ("标的", "item_name"),
        ("规格/要求", "specification"),
        ("数量", "quantity"),
        ("预算", "budget"),
        ("进口产品", "import_allowed"),
        ("页码", "page"),
    )
    return "；".join(
        f"{label}：{item.get(key)}"
        for label, key in labels
        if item.get(key) not in (None, "")
    )



def create_report(data):

    data = data if isinstance(data, dict) else {}


    doc = Document()


    # ======================
    # 标题
    # ======================

    title = doc.add_heading(
        "AI投标风险分析报告",
        level=1
    )


    doc.add_paragraph(
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )



    # ======================
    # 一、项目基本信息
    # ======================


    doc.add_heading(
        "一、项目基本信息",
        level=2
    )


    info = [

        ("项目名称",
        data.get("project_name") or "未找到"),

        ("招标单位",
        data.get("tender_company") or "未找到"),

        ("投标截止",
        data.get("deadline") or "未找到"),

        ("投标保证金",
        data.get("deposit") or "未找到")

    ]


    for k,v in info:

        doc.add_paragraph(
            f"{k}：{v}"
        )



    # ======================
    # 二、风险评分
    # ======================


    doc.add_heading(
        "二、投标风险评估",
        level=2
    )


    risks = [
        risk
        for risk in (data.get("risk") or [])
        if isinstance(risk, dict)
    ] if isinstance(data.get("risk") or [], list) else []
    summary = calculate_risk_summary(risks)



    doc.add_paragraph(
        f"投标建议指数：{summary['score']}分"
    )


    doc.add_paragraph(
        f"风险评级：{summary['score_level']}"
    )


    doc.add_paragraph(
        f"风险统计：共 {summary['risk_count']}项，"
        f"高风险 {summary['high_count']}项，"
        f"中风险 {summary['middle_count']}项，"
        f"低风险 {summary['low_count']}项"
    )


    doc.add_paragraph(
        f"风险总扣分：-{summary['total_deduction']}分"
    )



    # ======================
    # 三、技术要求
    # ======================


    doc.add_heading(
        "三、技术要求摘要",
        level=2
    )


    for item in data.get(
        "technical_requirements",
        []
    ):

        doc.add_paragraph(
            f"• {item}"
        )



    # ======================
    # 四、商务要求
    # ======================


    doc.add_heading(
        "四、商务资格要求",
        level=2
    )


    for item in data.get(
        "business_requirements",
        []
    ):

        doc.add_paragraph(
            f"• {item}"
        )



    # ======================
    # 五、采购清单
    # ======================

    procurement_items = data.get("procurement_requirements", [])
    has_procurement_section = isinstance(procurement_items, list) and bool(procurement_items)
    if has_procurement_section:
        doc.add_heading("五、采购清单摘要", level=2)
        for item in procurement_items:
            content = format_procurement_item(item)
            if content:
                doc.add_paragraph(f"• {content}")


    # ======================
    # 可选采购清单之后，剩余章节保持连续编号
    # ======================


    doc.add_heading(
        "六、风险明细" if has_procurement_section else "五、风险明细",
        level=2
    )


    table = doc.add_table(
        rows=1,
        cols=5
    )


    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )


    table.style="Table Grid"


    headers=[

        "等级",
        "页码",
        "风险",
        "风险原因",
        "处理建议"

    ]


    for i,h in enumerate(headers):

        table.rows[0].cells[i].text=h



    for risk in risks:


        row=table.add_row().cells


        row[0].text = (
            risk.get("level","")
        )


        report_page = risk.get("report_page") or risk.get("page", "")
        row[1].text = f"第{report_page}页"


        row[2].text = (
            risk.get("keyword","")
        )


        row[3].text = (
            risk.get("reason","")
        )


        row[4].text = (
            risk.get("suggestion","")
        )



    # ======================
    # 原文依据
    # ======================


    doc.add_heading(
        "七、风险原文依据" if has_procurement_section else "六、风险原文依据",
        level=2
    )


    for i,risk in enumerate(risks):


        doc.add_paragraph(
            f"风险{i+1}：{risk.get('keyword','')}"
        )


        doc.add_paragraph(
            risk.get(
                "quote",
                ""
            )
        )



    # ======================
    # AI建议
    # ======================


    doc.add_heading(
        "八、AI投标建议" if has_procurement_section else "七、AI投标建议",
        level=2
    )


    doc.add_paragraph(data.get("suggestion") or "暂无")



    # ======================
    # 保存
    # ======================


    report_dir = get_report_dir()
    report_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    path = os.path.join(
        report_dir,
        f"AI投标风险分析报告_{timestamp}.docx",
    )


    doc.save(path)


    return os.path.abspath(path)
