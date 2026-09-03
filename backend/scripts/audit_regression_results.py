import contextlib
import argparse
import io
import json
import os
import re
import sys
import tempfile
import shutil
from pathlib import Path

import fitz
from docx import Document


BACKEND_DIR = Path(__file__).resolve().parents[1]
PROJECT_ROOT = BACKEND_DIR.parent
SAMPLE_ROOT = PROJECT_ROOT / "regression_samples"


def compact(value):
    return re.sub(r"\s+", "", str(value or ""))


def document_text(path):
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def load_rendered_pages(sample_dir):
    sources = [
        path
        for path in sample_dir.iterdir()
        if path.is_file() and path.name.startswith("source.")
    ]
    if len(sources) != 1:
        raise RuntimeError(f"{sample_dir.name} must have exactly one source")
    source_path = sources[0]
    renderer = None

    if source_path.suffix.lower() in {".doc", ".docx"}:
        sys.path.insert(0, str(BACKEND_DIR))
        from app.services.word_preview import (
            ensure_word_preview_pdf,
            get_preview_renderer,
        )

        with tempfile.TemporaryDirectory(prefix="ai-bid-word-audit-") as temp_dir:
            temporary_source = Path(temp_dir) / source_path.name
            shutil.copy2(source_path, temporary_source)
            with contextlib.redirect_stdout(io.StringIO()):
                rendered_path = ensure_word_preview_pdf(str(temporary_source))
            renderer = get_preview_renderer(rendered_path)
            with fitz.open(rendered_path) as document:
                page_texts = [page.get_text() for page in document]
        return page_texts, source_path.suffix.lower().lstrip("."), renderer

    with fitz.open(source_path) as document:
        page_texts = [page.get_text() for page in document]
    return page_texts, source_path.suffix.lower().lstrip("."), renderer


def audit_sample(sample_dir):
    actual_path = sample_dir / "actual.local.json"
    actual = json.loads(actual_path.read_text(encoding="utf-8"))
    payload = actual.get("result") or {}
    risks = [risk for risk in (payload.get("risk") or []) if isinstance(risk, dict)]

    raw_page_texts, source_format, renderer = load_rendered_pages(sample_dir)
    page_texts = [compact(text) for text in raw_page_texts]

    sys.path.insert(0, str(BACKEND_DIR))
    from app.services.risk_location import ensure_locatable_highlights

    ensure_locatable_highlights(
        risks,
        [
            {"page": index + 1, "text": page_text}
            for index, page_text in enumerate(raw_page_texts)
        ],
    )

    risk_checks = []
    for index, risk in enumerate(risks, 1):
        try:
            page_number = int(risk.get("page"))
        except (TypeError, ValueError, OverflowError):
            page_number = 0
        page_valid = 1 <= page_number <= len(page_texts)
        page_text = page_texts[page_number - 1] if page_valid else ""
        keyword = compact(risk.get("keyword"))
        quote = compact(risk.get("quote"))
        highlight_words = risk.get("highlight_words") or []
        if not isinstance(highlight_words, list):
            highlight_words = []
        normalized_words = [str(word).strip() for word in highlight_words if str(word).strip()]
        raw_page_text = raw_page_texts[page_number - 1] if page_valid else ""

        risk_checks.append(
            {
                "case_id": f"RISK-{index:03d}",
                "page": page_number,
                "page_valid": page_valid,
                "keyword_present": bool(keyword),
                "keyword_page_match": bool(keyword and keyword in page_text),
                "quote_present": bool(quote),
                "quote_page_match": bool(quote and quote in page_text),
                "highlight_word_count": len(normalized_words),
                "highlight_page_match": any(
                    word in raw_page_text for word in normalized_words
                ),
            }
        )

    from app.services.report import create_report

    report_data = dict(payload.get("analysis") or {})
    report_data["project_name"] = payload.get("project_name")
    report_data["risk"] = risks
    with tempfile.TemporaryDirectory(prefix="ai-bid-report-audit-") as temp_dir:
        previous_dir = Path.cwd()
        os.chdir(temp_dir)
        try:
            with contextlib.redirect_stdout(io.StringIO()):
                report_path = create_report(report_data)
            report_text = document_text(report_path)
        finally:
            os.chdir(previous_dir)

    report_checks = {
        "opens": True,
        "project_name_match": str(payload.get("project_name") or "未找到") in report_text,
        "score_match": f"投标建议指数：{payload.get('score', 0)}分" in report_text,
        "risk_count_match": f"共 {payload.get('risk_count', 0)}项" in report_text,
        "deduction_match": f"风险总扣分：-{payload.get('total_deduction', 0)}分" in report_text,
        "procurement_items_match": all(
            str(item.get("item_name") or "") in report_text
            for item in (payload.get("analysis") or {}).get(
                "procurement_requirements", []
            )
            if isinstance(item, dict) and item.get("item_name")
        ),
    }

    summary = {
        "sample_id": actual.get("sample_id") or sample_dir.name,
        "source_format": source_format,
        "renderer": renderer,
        "page_count": len(page_texts),
        "risk_count": len(risks),
        "valid_pages": sum(check["page_valid"] for check in risk_checks),
        "keyword_page_matches": sum(
            check["keyword_page_match"] for check in risk_checks
        ),
        "quote_page_matches": sum(check["quote_page_match"] for check in risk_checks),
        "highlight_page_matches": sum(
            check["highlight_page_match"] for check in risk_checks
        ),
        "report_checks": report_checks,
        "risk_checks": risk_checks,
    }
    (sample_dir / "audit.local.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return {key: value for key, value in summary.items() if key != "risk_checks"}


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--sample-id", action="append", dest="sample_ids")
    return parser.parse_args()


def run():
    args = parse_args()
    summaries = []
    for sample_dir in sorted(SAMPLE_ROOT.iterdir()):
        if args.sample_ids and sample_dir.name not in set(args.sample_ids):
            continue
        has_source = any(
            path.is_file() and path.name.startswith("source.")
            for path in sample_dir.iterdir()
        )
        if has_source and (sample_dir / "actual.local.json").exists():
            summaries.append(audit_sample(sample_dir))
    if not summaries:
        raise RuntimeError("No completed local PDF regression results")
    for summary in summaries:
        print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    run()
